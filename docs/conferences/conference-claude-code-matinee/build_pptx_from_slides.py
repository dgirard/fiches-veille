#!/usr/bin/env python3
"""build_pptx_from_slides.py — convertit le deck HTML « Matinée Claude Code » en .pptx natif.

    python3 build_pptx_from_slides.py <slides.html> <out.pptx>

Conversion composant-par-composant du vocabulaire du deck (Sharp Artisan) :
covers de session, slides punch, grilles de cards/stats/steps, tableaux t-table,
listes, agenda, frise du cycle (cyc-row), arborescence (tree), checks, quotes.
Les notes de speaker (.notes) deviennent des notes PowerPoint.
Nécessite : beautifulsoup4, python-pptx. Idempotent.
"""
from __future__ import annotations

import re
import sys

from bs4 import BeautifulSoup, NavigableString, Tag
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

# ── Design tokens ────────────────────────────────────────────────────────────
NOIR = RGBColor(0x11, 0x11, 0x11)
NOIR_ELEV = RGBColor(0x1A, 0x1A, 0x1A)
BLANC = RGBColor(0xF9, 0xF9, 0xF9)
BLANC_PUR = RGBColor(0xFF, 0xFF, 0xFF)
CUIVRE = RGBColor(0xC0, 0x80, 0x40)
CUIVRE_CLAIR = RGBColor(0xE0, 0xA0, 0x40)
CUIVRE_LUM = RGBColor(0xF0, 0xB0, 0x50)
CUIVRE_FONCE = RGBColor(0x80, 0x50, 0x00)
GRIS_DARK_2 = RGBColor(0xB8, 0xB8, 0xB8)   # texte secondaire sur fond noir
GRIS_LIGHT_2 = RGBColor(0x3A, 0x3A, 0x3A)  # texte secondaire sur fond clair
GRIS_LIGHT_3 = RGBColor(0x6B, 0x6B, 0x6B)
SUCCESS = RGBColor(0x50, 0xB0, 0x50)
DANGER = RGBColor(0xC0, 0x60, 0x40)
CARD_DARK = RGBColor(0x22, 0x22, 0x22)
CARD_LIGHT = RGBColor(0xF3, 0xF3, 0xF3)

F_DISPLAY = "Epilogue"
F_BODY = "Space Grotesk"
F_MONO = "JetBrains Mono"

# deck conçu en 1920×1080 → EMU (12192000×6858000)
SCALE = 12192000 / 1920


def E(px: float) -> Emu:
    return Emu(int(px * SCALE))


def clean(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()


# ── Briques pptx ─────────────────────────────────────────────────────────────
def add_box(slide, x, y, w, h, fill=None, line=None, line_w=1.0, round_=True):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        E(x), E(y), E(w), E(h))
    if round_:
        try:
            shp.adjustments[0] = 0.045
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def add_text(slide, x, y, w, h, runs_list, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, wrap=True, space_after=6):
    """runs_list : liste de paragraphes ; chaque paragraphe = liste de
    (texte, taille_pt, gras, couleur, police)."""
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for para in runs_list:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        for txt, size, bold, color, font in para:
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.name = font
    return tb


def rich_runs(node: Tag, size: float, color, font, accent, bold_default=False):
    """Inline d'un nœud HTML → liste de runs, en respectant strong et u-accent."""
    runs = []
    def walk(n, bold, col):
        for ch in n.children:
            if isinstance(ch, NavigableString):
                t = str(ch).replace(" ", " ")
                t = re.sub(r"\s+", " ", t)
                if t:
                    runs.append((t, size, bold, col, font))
            elif isinstance(ch, Tag):
                b = bold or ch.name in ("strong", "b")
                c = col
                cls = " ".join(ch.get("class", []))
                style = ch.get("style", "")
                if "u-accent" in cls:
                    c = accent
                if "cuivre" in style:
                    c = accent
                elif "color:#fff" in style or "color: #fff" in style:
                    c = BLANC_PUR
                if ch.name == "br":
                    runs.append((" ", size, bold, col, font))
                    continue
                walk(ch, b, c)
    walk(node, bold_default, color)
    # trim
    out = []
    for i, (t, s, b, c, f) in enumerate(runs):
        if not out and not t.strip():
            continue
        out.append((t, s, b, c, f))
    return out or [(clean(node.get_text()), size, bold_default, color, font)]


# ── Convertisseur ────────────────────────────────────────────────────────────
class Deck:
    def __init__(self, html_path):
        self.soup = BeautifulSoup(open(html_path, encoding="utf-8").read(), "lxml")
        self.prs = Presentation()
        self.prs.slide_width = Emu(12192000)
        self.prs.slide_height = Emu(6858000)
        self.blank = self.prs.slide_layouts[6]
        self.idx = 0

    # palette selon fond
    def pal(self, dark):
        return {
            "fg": BLANC_PUR if dark else NOIR,
            "fg2": GRIS_DARK_2 if dark else GRIS_LIGHT_2,
            "fg3": GRIS_DARK_2 if dark else GRIS_LIGHT_3,
            "accent": CUIVRE_LUM if dark else CUIVRE_FONCE,
            "card": CARD_DARK if dark else CARD_LIGHT,
        }

    def new_slide(self, dark):
        self.idx += 1
        s = self.prs.slides.add_slide(self.blank)
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = NOIR if dark else BLANC
        return s

    def notes_and_pagenum(self, slide, sec, dark):
        n = sec.find("div", class_="notes")
        if n:
            slide.notes_slide.notes_text_frame.text = clean(n.get_text())
        add_text(slide, 1840, 1030, 60, 36,
                 [[(str(self.idx), 12, False, self.pal(dark)["fg3"], F_MONO)]],
                 align=PP_ALIGN.RIGHT)

    # ── types de slide ──
    def build(self):
        for sec in self.soup.find_all("section", class_="slide"):
            cls = " ".join(sec.get("class", []))
            dark = "slide--dark" in cls
            if "punch" in cls:
                self.slide_punch(sec)
            elif sec.find("h1", class_="t-cover"):
                self.slide_cover(sec, dark)
            else:
                self.slide_standard(sec, dark)
        return self.prs

    def slide_punch(self, sec):
        s = self.new_slide(True)
        p = self.pal(True)
        act = sec.find(class_="punch-act")
        if act:
            add_text(s, 96, 120, 1728, 50,
                     [[(clean(act.get_text()).upper(), 15, True, CUIVRE_LUM, F_BODY)]])
        title = sec.find(class_="punch-title")
        if title:
            add_text(s, 96, 330, 1728, 420,
                     [rich_runs(title, 54, BLANC_PUR, F_DISPLAY, CUIVRE_LUM, True)],
                     anchor=MSO_ANCHOR.MIDDLE)
        sub = sec.find(class_="punch-sub")
        if sub:
            add_text(s, 96, 790, 1500, 200,
                     [rich_runs(sub, 19, GRIS_DARK_2, F_BODY, CUIVRE_LUM)])
        self.notes_and_pagenum(s, sec, True)

    def slide_cover(self, sec, dark):
        s = self.new_slide(dark)
        p = self.pal(dark)
        y = 90
        head = sec.find(class_="sess-head")
        if head:
            num = head.find(class_="sess-num")
            meta = head.find(class_="sess-meta")
            if num:
                add_text(s, 96, 70, 300, 160,
                         [[(clean(num.get_text()), 96, True, CUIVRE_LUM, F_DISPLAY)]])
            if meta:
                add_text(s, 420, 130, 1300, 60,
                         [rich_runs(meta, 20, p["fg2"], F_BODY, p["accent"])])
            y = 300
        for eb in sec.find_all("p", class_="t-eyebrow", recursive=True):
            if eb.find_parent(class_="sess-head"):
                continue
            add_text(s, 96, 60, 1728, 44,
                     [[(clean(eb.get_text()).upper(), 14, True, p["accent"], F_BODY)]])
            y = max(y, 160)
            break
        title = sec.find("h1", class_="t-cover")
        if title:
            add_text(s, 96, y, 1728, 320,
                     [rich_runs(title, 60, p["fg"], F_DISPLAY, CUIVRE_LUM, True)])
            y += 350
        for q in sec.find_all("p", class_="quote"):
            add_text(s, 96, y, 1500, 170,
                     [rich_runs(q, 20, p["fg2"], F_BODY, p["accent"])])
            y += 180
        spk = sec.find(class_="sess-speaker")
        if spk:
            lines = [[(clean(x.get_text()), 13, False, p["fg3"], F_BODY)]
                     for x in spk.find_all("p")]
            add_text(s, 96, 960, 1500, 90, lines, space_after=2)
        else:
            extra = [c for c in sec.find_all("p", recursive=False)
                     if not set(c.get("class", [])) & {"quote", "t-eyebrow"}]
            yy = y
            for c in extra:
                add_text(s, 96, yy, 1600, 80,
                         [rich_runs(c, 17, p["fg2"], F_BODY, p["accent"])])
                yy += 70
        self.notes_and_pagenum(s, sec, dark)

    # ── slide standard : walker de composants ──
    def slide_standard(self, sec, dark):
        s = self.new_slide(dark)
        p = self.pal(dark)
        hdr = sec.find(class_="slide-header")
        if hdr:
            add_text(s, 96, 48, 1728, 40,
                     [[(clean(hdr.get_text()).upper(), 14, True, p["accent"], F_BODY)]])
        title = sec.find(class_="slide-title")
        y = 110.0
        if title:
            add_text(s, 96, y, 1728, 90,
                     [rich_runs(title, 34, p["fg"], F_DISPLAY, p["accent"], True)])
            y = 230.0
        body = self.collect_blocks(sec)
        # hauteur restante pour les blocs "extensibles"
        bottom = 1020.0
        fixed_h = sum(self.block_height(b) for b in body if not self.expands(b))
        n_exp = sum(1 for b in body if self.expands(b))
        free = max(120.0, bottom - y - fixed_h - 18 * max(0, len(body) - 1))
        exp_h = free / n_exp if n_exp else 0
        for b in body:
            h = exp_h if self.expands(b) else self.block_height(b)
            self.render_block(s, b, 96, y, 1728, h, dark)
            y += h + 18
        self.notes_and_pagenum(s, sec, dark)

    def collect_blocks(self, root):
        """Aplati la section en blocs de composants connus, dans l'ordre."""
        SKIP = {"slide-header", "notes", "slide-page-number"}
        blocks = []

        def walk(node):
            for ch in node.children:
                if not isinstance(ch, Tag):
                    continue
                cls = set(ch.get("class", []))
                if cls & SKIP or ch.name in ("h1", "h2") :
                    continue
                if "grid" in cls or "steps" in cls:
                    blocks.append(("grid", ch)); continue
                if ch.name == "table":
                    blocks.append(("table", ch)); continue
                if "cyc-row" in cls:
                    blocks.append(("cycle", ch)); continue
                if "cyc-legend" in cls:
                    blocks.append(("legend", ch)); continue
                if "tree" in cls:
                    blocks.append(("tree", ch)); continue
                if "agenda-row" in cls:
                    if blocks and blocks[-1][0] == "agenda":
                        blocks[-1][1].append(ch)
                    else:
                        blocks.append(("agenda", [ch]))
                    continue
                if "check" in cls and ch.name == "div":
                    if blocks and blocks[-1][0] == "checks":
                        blocks[-1][1].append(ch)
                    else:
                        blocks.append(("checks", [ch]))
                    continue
                if ch.name == "ul":
                    blocks.append(("list", ch)); continue
                if ch.name == "p":
                    blocks.append(("para", ch)); continue
                if ch.name == "div":
                    # stat isolé (layout 2 colonnes) ou wrapper → récursion
                    if "stat" in cls:
                        blocks.append(("grid", ch.parent if False else ch))
                        continue
                    walk(ch)

        walk(root)
        return blocks

    def expands(self, b):
        return b[0] in ("grid", "table", "cycle", "tree")

    def block_height(self, b):
        kind = b[1] if isinstance(b[1], list) else None
        t = b[0]
        if t == "para":
            return 64
        if t == "list":
            return 46 * len(b[1].find_all("li"))
        if t == "agenda":
            return 44 * len(b[1])
        if t == "checks":
            return 60
        if t == "legend":
            return 56
        return 200

    # ── rendu des blocs ──
    def render_block(self, s, b, x, y, w, h, dark):
        t, node = b
        p = self.pal(dark)
        if t == "para":
            cls = " ".join(node.get("class", []))
            size = 18 if "t-body-lg" in cls or "quote" in cls else 15
            if "quote" in cls:
                add_text(s, x, y, w, h, [rich_runs(node, 20, p["accent"], F_DISPLAY, p["accent"], True)])
            else:
                col = p["fg2"] if "t-muted" in cls else p["fg"]
                add_text(s, x, y, w, h, [rich_runs(node, size, col, F_BODY, p["accent"])])
        elif t == "list":
            paras = []
            for li in node.find_all("li"):
                runs = [("•  ", 16, True, p["accent"], F_BODY)]
                runs += rich_runs(li, 16, p["fg"], F_BODY, p["accent"])
                paras.append(runs)
            add_text(s, x, y, w, h, paras, space_after=10)
        elif t == "agenda":
            paras = []
            for row in node:
                hh = row.find(class_="agenda-h"); tt = row.find(class_="agenda-t"); aa = row.find(class_="agenda-a")
                paras.append([
                    ((clean(hh.get_text()) + "   ") if hh else "", 15, True, p["accent"], F_MONO),
                    (clean(tt.get_text()) if tt else "", 17, True, p["fg"], F_DISPLAY),
                    (("   — " + clean(aa.get_text())) if aa else "", 14, False, p["fg3"], F_BODY),
                ])
            add_text(s, x, y, w, h, paras, space_after=8)
        elif t == "checks":
            runs = []
            for c in node:
                runs.append(("✓ ", 20, True, SUCCESS, F_BODY))
                runs.append((clean(c.get_text()).lstrip("✓ ") + "    ", 20, True, p["fg"], F_DISPLAY))
            add_text(s, x, y, w, h, [runs])
        elif t == "tree":
            txt = node.get_text("\n") if "\n" not in node.get_text() else node.get_text()
            lines = [l.rstrip() for l in node.get_text().split("\n") if l.strip()]
            box = add_box(s, x, y, min(w, 900), h, fill=p["card"],
                          line=CUIVRE if not dark else CUIVRE_CLAIR)
            paras = [[(l, 15, False, p["fg"], F_MONO)] for l in lines]
            add_text(s, x + 36, y + 28, min(w, 900) - 72, h - 56, paras, space_after=3)
        elif t == "legend":
            runs = []
            for key in node.find_all(class_="cyc-key"):
                runs.append(("■ ", 13, True, CUIVRE_CLAIR, F_BODY))
                runs.append((clean(key.get_text()) + "      ", 13, False, p["fg2"], F_BODY))
            add_text(s, x, y, w, h, [runs])
        elif t == "table":
            self.render_table(s, node, x, y, w, h, dark)
        elif t == "cycle":
            self.render_cycle(s, node, x, y, w, h, dark)
        elif t == "grid":
            self.render_grid(s, node, x, y, w, h, dark)

    def render_table(self, s, node, x, y, w, h, dark):
        rows = node.find_all("tr")
        ncols = max(len(r.find_all(["td", "th"])) for r in rows)
        shp = s.shapes.add_table(len(rows), ncols, E(x), E(y), E(w), E(h))
        tbl = shp.table
        p = self.pal(dark)
        for i, r in enumerate(rows):
            for j, cell in enumerate(r.find_all(["td", "th"])):
                c = tbl.cell(i, j)
                c.fill.solid()
                c.fill.fore_color.rgb = (NOIR_ELEV if dark else CARD_LIGHT) if i else (CUIVRE_FONCE if not dark else NOIR)
                tf = c.text_frame
                tf.word_wrap = True
                para = tf.paragraphs[0]
                runs = rich_runs(cell, 13 if i else 12,
                                 p["fg"] if i else (BLANC_PUR if not dark else CUIVRE_LUM),
                                 F_BODY, p["accent"], bold_default=(i == 0))
                for txt, size, bold, col, font in runs:
                    rr = para.add_run(); rr.text = txt
                    rr.font.size = Pt(size); rr.font.bold = bold
                    rr.font.color.rgb = col; rr.font.name = font

    def render_cycle(self, s, node, x, y, w, h, dark):
        cells = node.find_all(class_="cyc-cell")
        n = len(cells)
        gap = 8
        cw = (w - gap * (n - 1)) / n
        ch = min(h, 230)
        cy = y + (h - ch) / 2
        for i, cell in enumerate(cells):
            cls = set(cell.get("class", []))
            cx = x + i * (cw + gap)
            if "cyc-cell--gate" in cls:
                box = add_box(s, cx, cy, cw, ch, fill=RGBColor(0x3A, 0x2A, 0x14),
                              line=CUIVRE_CLAIR, line_w=2.0)
            elif "cyc-cell--compound" in cls:
                box = add_box(s, cx, cy, cw, ch, fill=RGBColor(0x20, 0x20, 0x20),
                              line=RGBColor(0x8A, 0x8A, 0x8A), line_w=1.0)
            else:
                box = add_box(s, cx, cy, cw, ch, fill=NOIR_ELEV,
                              line=RGBColor(0x44, 0x44, 0x44), line_w=1.0)
            badge = cell.find(class_="cyc-badge")
            if badge:
                gate = "cyc-badge--gate" in set(badge.get("class", []))
                bw = 86
                bb = add_box(s, cx + (cw - bw) / 2, cy - 16, bw, 30,
                             fill=CUIVRE_CLAIR if gate else RGBColor(0x33, 0x33, 0x33),
                             line=None if gate else RGBColor(0x8A, 0x8A, 0x8A))
                tfb = bb.text_frame
                tfb.word_wrap = False
                tfb.margin_left = tfb.margin_right = tfb.margin_top = tfb.margin_bottom = 0
                pb = tfb.paragraphs[0]; pb.alignment = PP_ALIGN.CENTER
                rb = pb.add_run(); rb.text = clean(badge.get_text()).upper()
                rb.font.size = Pt(9); rb.font.bold = True; rb.font.name = F_BODY
                rb.font.color.rgb = NOIR if gate else BLANC_PUR
            num = cell.find(class_="cyc-num"); name = cell.find(class_="cyc-name"); desc = cell.find(class_="cyc-desc")
            gate_cell = "cyc-cell--gate" in cls
            paras = []
            if num:
                paras.append([(clean(num.get_text()), 11, False, GRIS_DARK_2, F_MONO)])
            if name:
                paras.append([(clean(name.get_text()), 13.5, True,
                               CUIVRE_LUM if gate_cell else BLANC_PUR, F_DISPLAY)])
            if desc:
                paras.append([(clean(desc.get_text()), 10.5, False, GRIS_DARK_2, F_BODY)])
            add_text(s, cx + 6, cy + 26, cw - 12, ch - 36, paras,
                     align=PP_ALIGN.CENTER, space_after=4)

    def render_grid(self, s, node, x, y, w, h, dark):
        p = self.pal(dark)
        kids = [c for c in node.children if isinstance(c, Tag)]
        cards = [c for c in kids if set(c.get("class", [])) & {"card", "stat", "step"}]
        if not cards:                       # wrapper inattendu → texte brut
            add_text(s, x, y, w, h, [[(clean(node.get_text())[:400], 14, False, p["fg2"], F_BODY)]])
            return
        n = len(cards)
        per_row = 2 if ("grid--2" in " ".join(node.get("class", [])) and n >= 2) else min(n, 3 if n != 4 else 4)
        if n == 5:
            per_row = 5
        rows = [cards[i:i + per_row] for i in range(0, n, per_row)]
        gap = 24
        rh = (h - gap * (len(rows) - 1)) / len(rows)
        for ri, row in enumerate(rows):
            cw = (w - gap * (len(row) - 1)) / len(row)
            for ci, card in enumerate(row):
                cx = x + ci * (cw + gap)
                cy = y + ri * (rh + gap)
                cls = set(card.get("class", []))
                accent_card = "card--accent" in cls
                line = CUIVRE_CLAIR if accent_card else (
                    DANGER if "card--rule-danger" in cls else (
                        SUCCESS if "card--rule-success" in cls else
                        (RGBColor(0x44, 0x44, 0x44) if dark else RGBColor(0xDD, 0xD5, 0xCD))))
                add_box(s, cx, cy, cw, rh, fill=p["card"], line=line,
                        line_w=2.0 if accent_card else 1.0)
                paras = []
                if "stat" in cls:
                    nn = card.find(class_="n"); cc = card.find(class_="c"); dd = card.find(class_="d")
                    if nn: paras.append([(clean(nn.get_text()), 44, True, p["accent"], F_DISPLAY)])
                    if cc: paras.append([(clean(cc.get_text()).upper(), 12, True, p["fg"], F_BODY)])
                    if dd: paras.append(rich_runs(dd, 12, p["fg2"], F_BODY, p["accent"]))
                elif "step" in cls:
                    for selector, size, bold, col, font in (
                            ("step-num", 12, True, p["accent"], F_MONO),
                            ("step-name", 18, True, p["fg"], F_DISPLAY),
                            ("step-desc", 12, False, p["fg2"], F_BODY),
                            ("step-ref", 11, False, p["fg3"], F_BODY)):
                        el = card.find(class_=selector)
                        if el:
                            paras.append([(clean(el.get_text()), size, bold, col, font)])
                else:  # card
                    eb = card.find(class_="card-eyebrow")
                    if eb:
                        col = p["accent"]
                        style = eb.get("style", "")
                        if "danger" in style: col = DANGER
                        if "success" in style: col = SUCCESS
                        paras.append([(clean(eb.get_text()).upper(), 12, True, col, F_BODY)])
                    ct = card.find(class_="card-title")
                    if ct:
                        paras.append(rich_runs(ct, 17, p["fg"], F_DISPLAY, p["accent"], True))
                    for el in card.children:
                        if not isinstance(el, Tag):
                            continue
                        ecls = set(el.get("class", []))
                        if ecls & {"card-eyebrow", "card-title"}:
                            continue
                        if el.name == "ul":
                            for li in el.find_all("li"):
                                runs = [("•  ", 12, True, p["accent"], F_BODY)]
                                runs += rich_runs(li, 12, p["fg2"], F_BODY, p["accent"])
                                paras.append(runs)
                        elif el.name == "p":
                            paras.append(rich_runs(el, 12, p["fg2"], F_BODY, p["accent"]))
                add_text(s, cx + 24, cy + 20, cw - 48, rh - 40, paras, space_after=6)


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    deck = Deck(sys.argv[1])
    prs = deck.build()
    prs.save(sys.argv[2])
    print(f"OK — {deck.idx} slides → {sys.argv[2]}")


if __name__ == "__main__":
    main()
